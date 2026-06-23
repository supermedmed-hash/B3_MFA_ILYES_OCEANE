from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

prs = Presentation()

def add_slide(title_text, content_text, speaker_notes, image_path=None):
    # Determine layout
    if image_path and os.path.exists(image_path):
        slide_layout = prs.slide_layouts[3] # Title, Content, Image (Two Content)
    else:
        slide_layout = prs.slide_layouts[1] # Title and Content

    slide = prs.slides.add_slide(slide_layout)
    
    # Set dark background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(30, 30, 30)

    # Title
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 200, 255)
    title.text_frame.paragraphs[0].font.name = 'Segoe UI'

    # Content
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = content_text
    for paragraph in tf.paragraphs:
        paragraph.font.color.rgb = RGBColor(230, 230, 230)
        paragraph.font.name = 'Segoe UI'
        paragraph.font.size = Pt(24)

    # Add image if provided
    if image_path and os.path.exists(image_path):
        # We replace the second placeholder or just insert it
        # Actually layout 3 has 2 content placeholders
        pic_shape = slide.placeholders[2]
        # remove the placeholder and add the image at its position
        left, top, width, height = pic_shape.left, pic_shape.top, pic_shape.width, pic_shape.height
        slide.shapes.element.remove(pic_shape.element)
        slide.shapes.add_picture(image_path, left, top, width, height)

    # Add speaker notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = speaker_notes

# Title Slide (Slide 1)
slide_1_layout = prs.slide_layouts[0]
slide1 = prs.slides.add_slide(slide_1_layout)
background1 = slide1.background
background1.fill.solid()
background1.fill.fore_color.rgb = RGBColor(15, 15, 15)
title1 = slide1.shapes.title
title1.text = "Smart Office 2.0"
title1.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 200, 255)
subtitle1 = slide1.placeholders[1]
subtitle1.text = "Soutenance Orale - Équipe B3\nOcéane, Florian, Ilyes, Mehdi"
subtitle1.text_frame.paragraphs[0].font.color.rgb = RGBColor(200, 200, 200)

slide1.notes_slide.notes_text_frame.text = "[Océane]\nBonjour à tous et merci de nous recevoir. Nous sommes l'équipe en charge de la transformation numérique de Biotech Corp. Aujourd'hui, nous allons vous présenter le socle technologique du projet 'Smart Office 2.0'. L'entreprise passant de 50 à 200 collaborateurs avec une forte culture du télétravail, notre mission a été de concevoir une infrastructure à la fois scalable, hyper-sécurisée et résiliente. Nous allons vous démontrer comment nous avons répondu à ce cahier des charges exigeant."

# Slide 2
add_slide(
    "Le Choix Stratégique - L'Architecture Hybride",
    "• Hyper-croissance de l'entreprise\n• Besoin de souveraineté des données\n• Réduction du TCO (-74% sur 3 ans)\n• Application web sur Azure PaaS",
    "[Océane]\nFace à l'hyper-croissance de Biotech Corp, le 100% local n'était plus viable financièrement, et le 100% Cloud posait des problèmes de souveraineté des données de R&D. Nous avons donc opté pour une approche Hybride.\nComme vous le voyez sur ce schéma conceptuel, nous conservons nos serveurs critiques 'On-Premise' pour la sécurité. En revanche, l'application métier Web est déportée sur le Cloud Microsoft Azure via le service PaaS 'App Service'.",
    "docs/gestion-projet/hybrid_cloud_concept.png"
)

# Slide 3
add_slide(
    "Architecture Réseau & Connectivité",
    "• Tunnel VPN IPsec IKEv2\n• Segmentation stricte en VLANs (Employés, R&D)\n• Wi-Fi 802.1X relié à l'Active Directory\n• Haute Disponibilité",
    "[Ilyes]\nMerci Océane. Le cœur de notre hybridation repose sur un tunnel VPN IPsec IKEv2 monté entre notre pare-feu local pfSense et la passerelle Azure. C'est ce tunnel qui permet à l'application Cloud de requêter nos bases locales de manière transparente et sécurisée.\nCôté LAN, nous avons structuré le réseau en VLANs stricts : le VLAN 20 pour les employés, et un VLAN 30 ultra-isolé pour la R&D. Le réseau Wi-Fi d'entreprise s'appuie sur une authentification 802.1X."
)

# Slide 4
add_slide(
    "Sécurité & Modèle Zero Trust",
    "• Pare-feu pfSense en 'Default Deny'\n• Aucune communication Employés -> R&D\n• Accès distants VPN OpenVPN sécurisés\n• Réduction drastique de la surface d'attaque",
    "[Ilyes]\nNotre politique de sécurité est dictée par le modèle Zero Trust. Le pare-feu pfSense agit en 'Default Deny'. Par exemple, le réseau des employés ne peut techniquement pas atteindre le réseau de la R&D.\nDe plus, face aux risques de cyberattaques et aux exigences du télétravail, nous avons implémenté des accès distants sécurisés via OpenVPN pour nos collaborateurs, avec une authentification forte."
)

# Slide 5
add_slide(
    "L'Usine Logicielle (DevOps)",
    "• Application conteneurisée avec Docker\n• CI/CD avec GitHub Actions\n• Azure Container Registry (ACR)\n• Déploiement en Zero Downtime",
    "[Mehdi]\nPour soutenir cette infrastructure moderne, nous avons adopté une démarche DevOps complète. Notre application métier est conteneurisée avec Docker. Cela garantit une portabilité parfaite.\nLe déploiement est entièrement automatisé via des pipelines CI/CD sur GitHub Actions. À chaque commit, une nouvelle image est générée, poussée dans notre registre Azure (ACR), et déployée en 'Zero Downtime' sur l'App Service Azure.",
    "docs/gestion-projet/smart_office_infra.png"
)

# Slide 6
add_slide(
    "Stratégie Data & Sauvegarde",
    "• Base PostgreSQL (Transactionnel)\n• Base MongoDB (Logs IoT)\n• Règle de sauvegarde du 3-2-1\n• Externalisation sur Azure Blob Storage",
    "[Mehdi]\nCôté données, nous utilisons une approche polyglotte : PostgreSQL pour garantir l'intégrité de nos réservations, et MongoDB pour absorber le flux massif de logs de nos capteurs IoT.\nPour protéger ces données vitales, nous appliquons la règle de sauvegarde du 3-2-1 : 3 copies des données, sur 2 supports différents (NAS Synology en RAID et disques SSD locaux), dont 1 copie externalisée et immuable sur un stockage Cloud."
)

# Slide 7
add_slide(
    "Supervision & Observabilité",
    "• Zabbix pour l'infrastructure et SNMP\n• Grafana pour la data en temps réel\n• Alerting intelligent via Slack/Pager\n• Monitoring du tunnel IPsec",
    "[Florian]\nUne telle infrastructure nécessite un monitoring proactif. Nous avons déployé une stack de supervision double : Zabbix pour l'IT, et Grafana pour le management et la data.\nComme vous le voyez sur cette maquette, Zabbix scrute l'état du tunnel VPN, l'espace disque de l'AD, et la charge de nos VM. Si le tunnel VPN IPsec tombe, des alertes sont remontées automatiquement. Grafana est directement pluggé sur notre base MongoDB pour afficher en temps réel l'utilisation des locaux.",
    "docs/gestion-projet/supervision_dashboard.png"
)

# Slide 8
add_slide(
    "Résilience et Gestion des Incidents",
    "• Plan de Reprise d'Activité (PRA)\n• Analyse d'Impact (BIA)\n• Objectif RTO : 4 heures maximum\n• Processus structurés ITSM",
    "[Florian]\nEnfin, la technique ne suffit pas sans processus. Nous avons structuré la gestion des incidents (ITSM) et réalisé une Analyse d'Impact sur l'Activité (BIA).\nCe BIA a conduit à notre Plan de Reprise d'Activité (PRA). Par exemple, en cas de sinistre majeur sur notre baie serveur principale On-Premise, notre objectif de délai de reprise (RTO) est fixé à 4 heures. L'application Azure, elle, restera en ligne."
)

# Slide 9
add_slide(
    "Bilan & Démonstration",
    "• Cahier des charges respecté\n• Hybridation parfaitement maîtrisée\n• Infrastructure scalable à 200+ employés\n• Démonstration de la plateforme",
    "[Équipe]\nPour conclure, le projet Smart Office 2.0 répond à tous les enjeux de Biotech Corp : une base locale souveraine et sécurisée, propulsée par la flexibilité du Cloud pour absorber la croissance. Nous avons respecté le budget, sécurisé les accès via du Zero Trust, et automatisé les tâches chronophages avec du CI/CD.\nNous vous remercions pour votre attention et sommes maintenant à votre disposition pour la démonstration."
)

# Save
os.makedirs("docs/gestion-projet", exist_ok=True)
prs.save("docs/gestion-projet/Presentation_SmartOffice2_B3.pptx")
print("PowerPoint generated successfully at docs/gestion-projet/Presentation_SmartOffice2_B3.pptx")
